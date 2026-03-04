"""
One-off script to generate the sample "Delivery Days by Suburb" spreadsheet.

Run once:  python create_sample_data.py

Staff should replace or extend the data with real suburb/postcode information.
Wednesday deliveries are calculated automatically by the backend:
  • Normal Wednesday  → same area as Friday's route
  • Monday is a WA public holiday → that Wednesday covers Monday's route instead
Therefore there is no Wednesday column in the spreadsheet.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUTPUT = os.path.join(os.path.dirname(__file__), "data", "Delivery Days by Suburb.xlsx")

HEADERS = ["Suburb", "Postcode", "Monday", "Tuesday", "Thursday", "Friday"]

# fmt: off
SAMPLE_ROWS = [
    # Suburb              Postcode  Mon    Tue    Thu    Fri
    ("Armadale",          "6112",   False, False, True,  False),
    ("Bayswater",         "6053",   False, True,  False, False),
    ("Belmont",           "6104",   False, True,  True,  False),
    ("Cannington",        "6107",   False, True,  True,  False),
    ("Canning Vale",      "6155",   True,  False, False, False),
    ("Claremont",         "6010",   False, False, False, True ),
    ("Cockburn Central",  "6164",   True,  False, False, True ),
    ("Cottesloe",         "6011",   True,  False, False, True ),
    ("East Perth",        "6004",   False, True,  True,  False),
    ("Fremantle",         "6160",   True,  False, False, True ),
    ("Guildford",         "6055",   True,  False, False, False),
    ("Joondalup",         "6027",   False, True,  True,  False),
    ("Kalamunda",         "6076",   False, False, True,  False),
    ("Karrinyup",         "6018",   False, False, False, True ),
    ("Kelmscott",         "6111",   False, False, True,  False),
    ("Leederville",       "6007",   False, True,  False, False),
    ("Mandurah",          "6210",   True,  False, False, True ),
    ("Midland",           "6056",   True,  False, False, False),
    ("Mirrabooka",        "6061",   False, True,  False, False),
    ("Morley",            "6062",   False, True,  False, False),
    ("Mount Lawley",      "6050",   False, True,  False, False),
    ("Mundaring",         "6073",   True,  False, False, False),
    ("Nedlands",          "6009",   False, False, False, True ),
    ("Northbridge",       "6003",   False, True,  False, False),
    ("O'Connor",          "6163",   True,  False, False, True ),
    ("Perth CBD",         "6000",   False, True,  True,  False),
    ("Rockingham",        "6168",   True,  False, False, True ),
    ("Scarborough",       "6019",   False, False, False, True ),
    ("South Perth",       "6151",   True,  False, False, True ),
    ("Spearwood",         "6163",   True,  False, False, True ),
    ("Stirling",          "6021",   False, False, False, True ),
    ("Subiaco",           "6008",   False, True,  True,  False),
    ("Swan View",         "6056",   True,  False, False, False),
    ("Victoria Park",     "6100",   False, True,  True,  False),
    ("Wangara",           "6065",   False, True,  False, False),
    ("Wanneroo",          "6065",   False, True,  False, False),
    ("Welshpool",         "6106",   False, True,  True,  False),
    ("Wembley",           "6014",   False, False, False, True ),
    ("Willetton",         "6155",   True,  False, False, False),
    ("Yokine",            "6060",   False, True,  False, False),
]
# fmt: on

HEADER_FILL  = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT  = Font(bold=True, color="FFFFFF")
DAY_FILL     = PatternFill("solid", fgColor="EAF4FF")
TICK_FILL    = PatternFill("solid", fgColor="D4EDDA")
TICK_FONT    = Font(color="155724", bold=True)
THIN         = Side(style="thin", color="CCCCCC")
BORDER       = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER       = Alignment(horizontal="center", vertical="center")


def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Delivery Days"

    # Note row
    ws.append(["NOTE: Wednesday deliveries are calculated automatically — "
               "normally Friday's area; Monday's area when Monday is a public holiday."])
    ws.merge_cells("A1:F1")
    note_cell = ws["A1"]
    note_cell.font = Font(italic=True, color="6C757D", size=9)
    note_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 18

    # Header row
    ws.append(HEADERS)
    for col, cell in enumerate(ws[2], 1):
        cell.fill   = HEADER_FILL
        cell.font   = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER if col > 2 else Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    # Data rows
    for row_data in SAMPLE_ROWS:
        ws.append(row_data)
        row_num = ws.max_row
        for col_idx, cell in enumerate(ws[row_num], 1):
            cell.border = BORDER
            if col_idx <= 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                is_tick = bool(cell.value)
                if is_tick:
                    cell.value = "Y"
                    cell.fill  = TICK_FILL
                    cell.font  = TICK_FONT
                else:
                    cell.value = ""
                    cell.fill  = DAY_FILL
                cell.alignment = CENTER
        ws.row_dimensions[row_num].height = 16

    # Column widths
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    for col in ("C", "D", "E", "F"):
        ws.column_dimensions[col].width = 11

    # Freeze header
    ws.freeze_panes = "A3"

    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    wb.save(OUTPUT)
    print(f"Saved: {OUTPUT}  ({len(SAMPLE_ROWS)} suburbs)")


if __name__ == "__main__":
    build()
