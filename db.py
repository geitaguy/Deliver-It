"""
Local SQLite database for Deliver-It.

Track-POD is the source of truth for order/route data.
This database stores:
  • order_meta     — which orders have been seen/viewed in this app
  • suburbs        — suburb delivery-day schedule (imported from Excel)
  • route_overrides — operational overrides (cancelled, full, rescheduled dates)
"""

import sqlite3
import os
from contextlib import contextmanager

DB_PATH = os.environ.get("DATABASE_PATH", "deliver_it.db")

ROUTING_AREAS = {"Monday", "Tuesday", "Thursday", "Friday"}
VALID_STATUSES = {"cancelled", "full", "limited", "rescheduled"}


def _connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


@contextmanager
def get_db():
    conn = _connect()
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def init_db():
    """Create all tables if they don't exist."""
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS order_meta (
                order_number    TEXT PRIMARY KEY,
                first_seen_at   TEXT NOT NULL DEFAULT (datetime('now')),
                first_viewed_at TEXT,
                notes           TEXT
            )
        """)

        conn.execute("""
            CREATE TABLE IF NOT EXISTS suburbs (
                id       INTEGER PRIMARY KEY AUTOINCREMENT,
                suburb   TEXT    NOT NULL,
                postcode TEXT    NOT NULL DEFAULT '',
                monday   INTEGER NOT NULL DEFAULT 0,
                tuesday  INTEGER NOT NULL DEFAULT 0,
                thursday INTEGER NOT NULL DEFAULT 0,
                friday   INTEGER NOT NULL DEFAULT 0
            )
        """)
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_suburbs_suburb   ON suburbs (suburb)"
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_suburbs_postcode ON suburbs (postcode)"
        )

        conn.execute("""
            CREATE TABLE IF NOT EXISTS route_overrides (
                id               INTEGER PRIMARY KEY AUTOINCREMENT,
                date             TEXT    NOT NULL,
                routing_area     TEXT    NOT NULL,
                status           TEXT    NOT NULL DEFAULT 'cancelled',
                rescheduled_date TEXT,
                message          TEXT,
                created_at       TEXT    NOT NULL DEFAULT (datetime('now')),
                UNIQUE (date, routing_area)
            )
        """)


# ---------------------------------------------------------------------------
# order_meta helpers (unchanged from original)
# ---------------------------------------------------------------------------

def mark_orders_seen(order_numbers: list[str]):
    """Record that we fetched these orders from Track-POD."""
    with get_db() as conn:
        conn.executemany(
            """
            INSERT INTO order_meta (order_number, first_seen_at)
            VALUES (?, datetime('now'))
            ON CONFLICT(order_number) DO NOTHING
            """,
            [(n,) for n in order_numbers],
        )


def mark_order_viewed(order_number: str):
    """Record that the user opened / acknowledged this order in the UI."""
    with get_db() as conn:
        conn.execute(
            """
            INSERT INTO order_meta (order_number, first_seen_at, first_viewed_at)
            VALUES (?, datetime('now'), datetime('now'))
            ON CONFLICT(order_number) DO UPDATE
              SET first_viewed_at = COALESCE(first_viewed_at, datetime('now'))
            """,
            (order_number,),
        )


def get_viewed_set() -> set[str]:
    """Return order numbers that have been viewed in this app."""
    with get_db() as conn:
        rows = conn.execute(
            "SELECT order_number FROM order_meta WHERE first_viewed_at IS NOT NULL"
        ).fetchall()
    return {row["order_number"] for row in rows}


def get_meta(order_number: str) -> dict | None:
    with get_db() as conn:
        row = conn.execute(
            "SELECT * FROM order_meta WHERE order_number = ?", (order_number,)
        ).fetchone()
    return dict(row) if row else None


# ---------------------------------------------------------------------------
# Suburb helpers
# ---------------------------------------------------------------------------

def count_suburbs() -> int:
    with get_db() as conn:
        return conn.execute("SELECT COUNT(*) FROM suburbs").fetchone()[0]


def search_suburbs(query: str, limit: int = 15) -> list[dict]:
    """Case-insensitive substring search on suburb name or postcode."""
    q = f"%{query.strip()}%"
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT id, suburb, postcode, monday, tuesday, thursday, friday
            FROM   suburbs
            WHERE  suburb   LIKE ? OR postcode LIKE ?
            ORDER  BY suburb
            LIMIT  ?
            """,
            (q, q, limit),
        ).fetchall()
    return [dict(r) for r in rows]


def import_suburbs_from_excel(filepath: str) -> int:
    """
    Replace all suburb rows with data from the Excel spreadsheet.
    Returns the number of rows imported (0 on failure).

    Expected spreadsheet columns (header auto-detected):
        Suburb | Postcode | Monday | Tuesday | Thursday | Friday
    Wednesday is intentionally absent — it is derived dynamically.
    """
    if not os.path.exists(filepath):
        return 0

    try:
        import openpyxl
    except ImportError:
        return 0

    _DAY_ORDER_ALL = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

    def _is_truthy(v: object) -> bool:
        if v is None:
            return False
        if isinstance(v, bool):
            return v
        if isinstance(v, (int, float)):
            return bool(v)
        return str(v).strip().lower() not in ("", "false", "no", "0", "n", "none")

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active

    suburb_col: int | None  = None
    postcode_col: int | None = None
    day_cols: dict[str, int] = {}
    data_start_row = 2

    for ri, row in enumerate(ws.iter_rows(max_row=6, values_only=True), 1):
        for ci, val in enumerate(row, 1):
            key = str(val or "").strip().title()
            if key in ("Suburb", "Suburb Name"):
                suburb_col = ci
            elif key in ("Postcode", "Post Code", "Zip"):
                postcode_col = ci
            elif key in _DAY_ORDER_ALL:
                day_cols[key] = ci
        if day_cols:
            data_start_row = ri + 1
            break

    if not day_cols:
        wb.close()
        return 0

    rows: list[tuple] = []
    for row in ws.iter_rows(min_row=data_start_row, values_only=True):
        suburb   = str(row[suburb_col   - 1] or "").strip() if suburb_col   else ""
        postcode = str(row[postcode_col - 1] or "").strip() if postcode_col else ""
        if not suburb and not postcode:
            continue
        if postcode.endswith(".0"):
            postcode = postcode[:-2]

        rows.append((
            suburb,
            postcode,
            int(_is_truthy(row[day_cols["Monday"]   - 1])) if "Monday"   in day_cols else 0,
            int(_is_truthy(row[day_cols["Tuesday"]  - 1])) if "Tuesday"  in day_cols else 0,
            int(_is_truthy(row[day_cols["Thursday"] - 1])) if "Thursday" in day_cols else 0,
            int(_is_truthy(row[day_cols["Friday"]   - 1])) if "Friday"   in day_cols else 0,
        ))

    wb.close()

    if not rows:
        return 0

    with get_db() as conn:
        conn.execute("DELETE FROM suburbs")
        conn.executemany(
            "INSERT INTO suburbs (suburb, postcode, monday, tuesday, thursday, friday) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            rows,
        )

    return len(rows)


# ---------------------------------------------------------------------------
# Route override helpers
# ---------------------------------------------------------------------------

def get_route_overrides(
    from_date: str | None = None,
    to_date: str | None = None,
) -> list[dict]:
    """Return overrides, optionally filtered to a date range."""
    with get_db() as conn:
        if from_date and to_date:
            rows = conn.execute(
                "SELECT * FROM route_overrides "
                "WHERE date BETWEEN ? AND ? ORDER BY date, routing_area",
                (from_date, to_date),
            ).fetchall()
        elif from_date:
            rows = conn.execute(
                "SELECT * FROM route_overrides "
                "WHERE date >= ? ORDER BY date, routing_area",
                (from_date,),
            ).fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM route_overrides ORDER BY date, routing_area"
            ).fetchall()
    return [dict(r) for r in rows]


def upsert_route_override(
    date: str,
    routing_area: str,
    status: str,
    rescheduled_date: str | None,
    message: str | None,
) -> dict:
    """Insert or replace an override for (date, routing_area). Returns the saved row."""
    with get_db() as conn:
        conn.execute(
            """
            INSERT INTO route_overrides (date, routing_area, status, rescheduled_date, message)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(date, routing_area) DO UPDATE SET
                status           = excluded.status,
                rescheduled_date = excluded.rescheduled_date,
                message          = excluded.message,
                created_at       = created_at
            """,
            (date, routing_area, status, rescheduled_date, message),
        )
        row = conn.execute(
            "SELECT * FROM route_overrides WHERE date = ? AND routing_area = ?",
            (date, routing_area),
        ).fetchone()
    return dict(row)


def delete_route_override(override_id: int) -> bool:
    with get_db() as conn:
        cur = conn.execute(
            "DELETE FROM route_overrides WHERE id = ?", (override_id,)
        )
    return cur.rowcount > 0
